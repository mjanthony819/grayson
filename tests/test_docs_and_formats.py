"""v0.2 formats: documents, extra image/audio/video targets, crop, merge."""
import shutil
import subprocess

import pytest
from PIL import Image, features

from grayson.cli import run_action
from grayson.core import ACTIONS, actions_for, kind_of
from grayson.convert import docs

needs_office = pytest.mark.skipif(
    shutil.which("soffice") is None and shutil.which("libreoffice") is None,
    reason="LibreOffice not installed")
needs_ffmpeg = pytest.mark.skipif(shutil.which("ffmpeg") is None, reason="no ffmpeg")


def _png(p, size=(400, 300)):
    Image.new("RGBA", size, (200, 30, 30, 255)).save(p)
    return p


def test_kinds():
    for name, kind in [("a.docx", "word"), ("a.rtf", "word"), ("a.pptx", "slides"),
                       ("a.xlsx", "sheet"), ("a.csv", "sheet"), ("a.md", "text"),
                       ("a.heic", "image"), ("a.opus", "audio"), ("a.wmv", "video")]:
        assert kind_of(name) == kind


def test_every_kind_has_actions_and_ids_unique():
    for f in ["a.png", "a.mp4", "a.mp3", "a.pdf", "a.docx", "a.pptx", "a.xlsx", "a.txt"]:
        assert actions_for([f]), f
    assert len({a.id for a in ACTIONS}) == len(ACTIONS)


def test_merge_hidden_for_single_file_shown_for_two():
    assert "auds-merge" not in {a.id for a in actions_for(["a.mp3"])}
    assert "auds-merge" in {a.id for a in actions_for(["a.mp3", "b.mp3"])}


def test_mixed_kinds_offer_nothing():
    assert actions_for(["a.docx", "b.pdf"]) == []


@pytest.mark.parametrize("action,ext", [
    ("img-to-bmp", ".bmp"), ("img-to-tiff", ".tiff"), ("img-to-gif", ".gif"),
    ("img-to-ico", ".ico"), ("img-to-heic", ".heic"), ("img-rotate", ".png"),
])
def test_image_targets(tmp_path, action, ext):
    src = _png(tmp_path / "a.png")
    out = run_action(action, [src])[0]
    assert out.suffix == ext and out.stat().st_size > 0


@pytest.mark.skipif(not features.check("avif"), reason="Pillow built without AVIF")
def test_avif(tmp_path):
    out = run_action("img-to-avif", [_png(tmp_path / "a.png")])[0]
    assert out.suffix == ".avif"


def test_resize_keeps_source_format(tmp_path):
    """Regression: v0.1 wrote PNG bytes into a .jpg after resize."""
    Image.new("RGB", (200, 100), "blue").save(tmp_path / "a.jpg")
    out = run_action("img-resize-50", [tmp_path / "a.jpg"])[0]
    with Image.open(out) as im:
        assert im.format == "JPEG" and im.size == (100, 50)


def test_crops(tmp_path):
    src = _png(tmp_path / "a.png")  # 400x300
    with Image.open(run_action("img-crop-square", [src])[0]) as im:
        assert im.size == (300, 300)
    with Image.open(run_action("img-crop-16x9", [src])[0]) as im:
        assert im.size == (400, 225)
    out = run_action("img-crop", [src], {"left": 10, "top": 0, "right": 10, "bottom": 50})[0]
    with Image.open(out) as im:
        assert im.size == (320, 150)
    with pytest.raises(ValueError):
        run_action("img-crop", [src], {"left": 60, "right": 60})


def test_ico_has_standard_sizes(tmp_path):
    out = run_action("img-to-ico", [_png(tmp_path / "a.png")])[0]
    with Image.open(out) as im:
        assert (256, 256) in im.info["sizes"] or (128, 128) in im.info["sizes"]


def test_images_to_docx(tmp_path):
    from docx import Document
    a, b = _png(tmp_path / "a.png"), _png(tmp_path / "b.png", (100, 200))
    out = run_action("imgs-to-docx", [a, b])[0]
    assert len(Document(out).inline_shapes) == 2


def test_text_to_pdf_builtin_no_office(tmp_path):
    import pymupdf
    (tmp_path / "n.md").write_text("# Title\n\nhello **world**\n")
    out = run_action("text-to-pdf", [tmp_path / "n.md"])[0]
    with pymupdf.open(out) as d:
        assert "hello world" in d[0].get_text()


def test_text_to_html_and_txt(tmp_path):
    (tmp_path / "n.md").write_text("# Title\n\n- a\n- b\n")
    assert "<h1>Title</h1>" in run_action("text-to-html", [tmp_path / "n.md"])[0].read_text()
    h = tmp_path / "p.html"
    h.write_text("<html><body><p>one</p><script>x()</script><p>two</p></body></html>")
    txt = run_action("text-to-txt", [h])[0].read_text()
    assert "one" in txt and "two" in txt and "x()" not in txt


def test_pdf_to_txt_docx_jpg(tmp_path):
    from docx import Document
    (tmp_path / "n.txt").write_text("alpha beta")
    pdf = run_action("text-to-pdf", [tmp_path / "n.txt"])[0]
    assert "alpha beta" in run_action("pdf-to-txt", [pdf])[0].read_text()
    docx = run_action("pdf-to-docx", [pdf])[0]
    assert "alpha beta" in " ".join(p.text for p in Document(docx).paragraphs)
    assert run_action("pdf-to-jpg", [pdf])[0].suffix == ".jpg"


@needs_office
def test_office_roundtrips(tmp_path):
    import pymupdf
    (tmp_path / "n.md").write_text("# Title\n\nhello world\n")
    docx = run_action("text-to-docx", [tmp_path / "n.md"])[0]
    pdf = run_action("word-to-pdf", [docx])[0]
    with pymupdf.open(pdf) as d:
        assert "hello world" in d[0].get_text()
    assert run_action("word-to-odt", [docx])[0].suffix == ".odt"
    pngs = run_action("word-to-png", [docx])
    assert pngs and pngs[0].suffix == ".png"
    (tmp_path / "c.csv").write_text("a,b\n1,2\n")
    xlsx = run_action("sheet-to-xlsx", [tmp_path / "c.csv"])[0]
    assert run_action("sheet-to-csv", [xlsx])[0].suffix == ".csv"
    assert run_action("sheet-to-pdf", [tmp_path / "c.csv"])[0].suffix == ".pdf"


@needs_office
def test_slides_to_pdf_and_png(tmp_path):
    pptx = pytest.importorskip("pptx")
    p = pptx.Presentation()
    p.slides.add_slide(p.slide_layouts[0]).shapes.title.text = "One"
    p.slides.add_slide(p.slide_layouts[0]).shapes.title.text = "Two"
    p.save(tmp_path / "d.pptx")
    assert len(run_action("slides-to-png", [tmp_path / "d.pptx"])) == 2
    assert run_action("slides-to-pdf", [tmp_path / "d.pptx"])[0].suffix == ".pdf"


def test_missing_office_message(monkeypatch):
    monkeypatch.setenv("PATH", "")
    monkeypatch.delenv("GRAYSON_SOFFICE", raising=False)
    monkeypatch.delenv("ProgramFiles", raising=False)
    monkeypatch.delenv("ProgramFiles(x86)", raising=False)
    monkeypatch.setattr(docs.shutil, "which", lambda *_: None)
    with pytest.raises(docs.OfficeNotFound, match="winget install"):
        docs.find_soffice()


@needs_ffmpeg
def test_audio_targets_and_merge(tmp_path):
    wav = tmp_path / "s.wav"
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i",
                    "sine=d=1", str(wav)], check=True)
    for action, ext in [("aud-to-flac", ".flac"), ("aud-to-m4a", ".m4a"),
                        ("aud-to-aac", ".aac"), ("aud-to-opus", ".opus"),
                        ("aud-compress", ".mp3")]:
        assert run_action(action, [wav])[0].suffix == ext
    merged = run_action("auds-merge", [wav, wav])[0]
    out = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration",
                          "-of", "csv=p=0", str(merged)], capture_output=True, text=True) \
        if shutil.which("ffprobe") else None
    if out:
        assert 1.8 < float(out.stdout) < 2.3


@needs_ffmpeg
def test_video_targets(tmp_path):
    mp4 = tmp_path / "v.mp4"
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i",
                    "testsrc=d=1:s=320x240", "-f", "lavfi", "-i", "sine=d=1",
                    "-shortest", str(mp4)], check=True)
    for action, ext in [("vid-to-mkv", ".mkv"), ("vid-to-mov", ".mov"),
                        ("vid-to-avi", ".avi"), ("vid-720p", ".mp4"),
                        ("vid-mute", ".mp4"), ("vid-extract-wav", ".wav")]:
        assert run_action(action, [mp4])[0].suffix == ext
